import numpy as np
import pandas as pd
from openpyxl import load_workbook


def strip_n_upper(df, columns):
    for col in columns:
        if df[col].dtype == 'object':  # Only apply to string columns
            df[col] = df[col].str.replace(r"\s+", " ", regex=True)
            df[col] = df[col].str.upper()
            df[col] = df[col].str.strip()

    return df


def clean_numeric_columns(df, columns):
    """To clean up numeric columns"""
    for col in columns:
        df[col] = (
            df[col]
            .astype(str)
            .str.strip()
            .str.replace("N", "")
            .str.replace(",", "")
            .str.replace("-", "")
            .str.replace("                                              -  ", "")
        )
        df[col] = pd.to_numeric(df[col], errors="coerce")
        df[col] = df[col].fillna(0)
        return df


# def clean_dates(date_value):
#     # Check if the value is an integer and greater than a specific threshold
#     if isinstance(date_value, int) and date_value > 0:
#         # Convert the Excel serial date to a datetime object
#         return pd.to_datetime(date_value, origin="1899-12-30", unit="D", dayfirst=False)
#     elif isinstance(date_value, str) and date_value.strip():
#         # Convert string date formats directly
#         return pd.to_datetime(str(date_value).strip(), dayfirst=False)
#     else:
#         return pd.to_datetime(
#             str(date_value).strip(), dayfirst=False
#         )  # .astype(str).str.strip())

def clean_date_column(column):
    """
    Clean and standardize a column with mixed date formats.
    Handles datetime objects, string dates, and numeric serial dates.
    """
    # print(f"Column sample: {column.head()}")
    # print(f"Column dtype: {column.dtype}")

    # Ensure consistent data type for processing
    column = column.astype(str).fillna("").infer_objects(copy=False).astype(str).str.strip()  # Handle NaN and convert to strings

    # Function to check if a value is an Excel serial number
    def is_excel_serial(value):
        try:
            return float(value).is_integer() and float(value) > 0
        except ValueError:
            return False

    # Apply checks for serial dates
    is_serial = column.apply(is_excel_serial)
    if is_serial.any():
        # print(f"Found {is_serial.sum()} Excel serial numbers. Converting...")
        column[is_serial] = pd.to_datetime(
            column[is_serial].astype(float),
            origin="1899-12-30",
            unit="D",
            errors="coerce",  # Convert invalid entries to NaT
        )

    # Handle string dates
    is_string_date = ~is_serial  # Non-serial values
    if is_string_date.any():
        # print(f"Found {is_string_date.sum()} string dates. Converting...")
        column[is_string_date] = pd.to_datetime(
            column[is_string_date],
            dayfirst=False,
            format="mixed",
            errors="coerce",  # Convert invalid string dates to NaT
        )

    # Fill NaT with a default value (like a specific date, or None, or NaN)
    column = column.astype("datetime64[ns]").fillna(pd.to_datetime("1900-01-01")).infer_objects(copy=False)  # or use None/NaN if required

    # Replace default "1900-01-01" with NaT
    column = column.replace(pd.to_datetime("1900-01-01"), pd.NaT)

    # Ensure the entire column is datetime
    return column










# Function to adjust column width dynamically
def auto_adjust_column_widths(excel_file):
    # Load the workbook
    workbook = load_workbook(excel_file)

    # Loop through all sheets in the workbook
    for sheet in workbook.sheetnames:
        worksheet = workbook[sheet]

        # Loop through all columns in the worksheet
        for col in worksheet.columns:
            max_length = 0
            column = col[
                0
            ].column_letter  # Get the column letter (e.g., "A", "B", etc.)
            for cell in col:
                try:
                    # Calculate the maximum length of the content in each column
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                except Exception:
                    pass
            # Set the column width based on the maximum length found
            adjusted_width = max_length + 1
            worksheet.column_dimensions[column].width = adjusted_width

    # Save the modified workbook
    workbook.save(excel_file)


def table_headers(df, header1, header2, split_header):
    if header2 != None:
        if split_header is None:
            higher_level_1 = [header1] * df.shape[1]  # First level of text
            higher_level_2 = [header2] * df.shape[1]  # Second level of text

            # Combine the higher-level labels with the existing column names
            df_copy = df.copy()
            df_copy.columns = pd.MultiIndex.from_tuples(
                zip(higher_level_1, higher_level_2, df.columns)
            )
            header_columns_df = pd.DataFrame(columns=df_copy.columns)

            return header_columns_df
        else:
            higher_level_1 = [header1] * df.shape[1]  # First level of text
            higher_level_2 = [header2] * df.shape[1]  # Second level of text
            higher_level_2[0] = split_header
            # Combine the higher-level labels with the existing column names
            df_copy = df.copy()
            df_copy.columns = pd.MultiIndex.from_tuples(
                zip(higher_level_1, higher_level_2, df.columns)
            )
            header_columns_df = pd.DataFrame(columns=df_copy.columns)
    else:
        if split_header is None:
            higher_level_1 = [header1] * df.shape[1]  # First level of text

            # Combine the higher-level labels with the existing column names
            df_copy = df.copy()
            df_copy.columns = pd.MultiIndex.from_tuples(
                zip(higher_level_1, df.columns)
            )
            header_columns_df = pd.DataFrame(columns=df_copy.columns)

            return header_columns_df
        else:
            higher_level_1 = [header1] * df.shape[1]  # First level of text
            higher_level_1[0] = split_header
            # Combine the higher-level labels with the existing column names
            df_copy = df.copy()
            df_copy.columns = pd.MultiIndex.from_tuples(
                zip(higher_level_1, df.columns)
            )
            header_columns_df = pd.DataFrame(columns=df_copy.columns)

    return header_columns_df

# Function to format a single DataFrame
def format_dataframe(df, percentage_cols, date_cols, numeric_but_text, num_over_perc):
    # Formatting function
    def format_value(x, col):
        if col in date_cols:
            return pd.to_datetime(x).date()
        elif col in numeric_but_text:
            return str(x) if str(x) != "nan" else np.nan
        elif str(x) == "nan":
            return np.nan
        elif isinstance(x, int):
            return f"{x:,}"  # Format integers with commas
        elif isinstance(x, float):
            if num_over_perc == "Y" and col in percentage_cols:
                return f"{x:,}"
            elif col in percentage_cols:
                return f"{x * 100:.2f}%"  # Format as percentage
            else:
                return f"{x:,.2f}"  # Format floats with commas and 2 decimal places
        else:
            return x  # Return the value as is for other types (e.g., strings)

    # Apply formatting function to each cell
    return df.apply(lambda col: col.apply(lambda x: format_value(x, col.name)))


# Remove extra labels before main data within files
def clean_input_sheet(df):
    # To get the point where the main data within an excel file starts
    def get_data_start(col):
        return "Y" if col == "S/N" else np.nan

    output = df.copy()

    if ("Unnamed: 0" in output.columns) | ("Unnamed: 1" in output.columns):
        output["data_start"] = output.iloc[:,0].apply(get_data_start)
        if output["data_start"].notna().sum() == 0:
            output["data_start"] = output.iloc[:,1].apply(get_data_start)
        col_headers = output[output["data_start"] == "Y"].iloc[:, :-1].values[0]
        col_headers = [
            str(col).upper().strip().replace("\n", " ") for col in col_headers
        ]
        output["data_start"] = output["data_start"].bfill()
        output = output[output["data_start"].isna()].iloc[:, :-1].reset_index(drop=True)

        output.columns = [str(col).lower().strip() for col in col_headers]
    else:
        col_headers = [
            str(col).upper().strip().replace("\n", " ") for col in output.columns
        ]
        output.columns = [str(col).lower().strip() for col in col_headers]
    
    if "account number" in output.columns:
        output = output[output["account number"].notna()]
    elif "account_number" in output.columns:
        output = output[output["account_number"].notna()]
    elif "a/c no." in output.columns:
        output = output[output["a/c no."].notna()]
    elif "acct_num" in output.columns:
        output = output[output["acct_num"].notna()]
    elif "cif id" in output.columns:
        output = output[output["cif id"].notna()]
    else:
        output = output[output["cifid"].notna()]

    output.columns = col_headers

    output = output.dropna(how="all", axis=1)
    if "NAN" in output.columns:
        output = output.drop(columns=["NAN"])

    return output


# # PD LGD MIGRATION FUNCTIONS
# def get_avg_lgd_migration_rate(df, avg_lgd_df):
#     # df = lgd_summary_per_segment_repay_adj.copy()
#     df_ = df.copy()
#     for cols in df_.columns:
#         if cols not in ["PERFORMING STATUS", "TOTAL"]:
#             df_.loc[:, cols] = df_.loc[:, cols] / df_.loc[:, "TOTAL"]

#     output = pd.concat([avg_lgd_df, df_], ignore_index=True)
#     return output


# def add_totals_and_identifiers(df, segment):
#     # Add Totals both to rows and columns
#     # Sum across Columns
#     df["TOTAL"] = df.sum(axis=1, numeric_only=True)

#     # Calculate the sum of numeric columns
#     total_row = df.select_dtypes(include="number").sum()
#     total_row["PERFORMING STATUS"] = "TOTAL"  # The label
#     total_row = pd.DataFrame([total_row])  # Convert to Dataframe

#     # Append the total row to the original DataFrame
#     df = pd.concat([df, total_row], ignore_index=True)

#     # Add segment Identifier
#     df["segment_identifier"] = segment

#     # Rearrange columns
#     df = df[
#         [
#             "segment_identifier",
#             "PERFORMING STATUS",
#             "CR",
#             "RR",
#             "SUBSTANDARD",
#             "DOUBTFUL",
#             "LOST",
#             "TOTAL",
#         ]
#     ]

#     # Rename Columns:
#     # df = df.rename(columns={"PERFORMING STATUS":""})
#     return df


# loan_class_of_interest = ["SUBSTANDARD", "DOUBTFUL", "LOST"]


# def calculate_final_lgd(row_index, df):
#     level_1 = df.at[0, "INITIAL LGD"]
#     level_2 = df.at[1, "INITIAL LGD"]
#     level_3 = df.at[2, "INITIAL LGD"]

#     if row_index == 0:
#         return (
#             level_1
#             if (level_1 <= level_2 and level_2 <= level_3)
#             else (
#                 level_1
#                 if (level_1 >= level_2 and level_1 >= level_3)
#                 else (
#                     level_1 if (level_1 >= level_2 and level_1 < level_3) else level_1
#                 )
#             )
#         )
#     elif row_index == 1:
#         return (
#             level_2
#             if (level_1 <= level_2 and level_2 <= level_3)
#             else (
#                 level_1
#                 if (level_1 >= level_2 and level_1 >= level_3)
#                 else (
#                     level_1
#                     if (level_1 >= level_2 and level_1 < level_3)
#                     else (
#                         level_2
#                         if (level_1 <= level_2 and level_2 >= level_3)
#                         else level_2
#                     )
#                 )
#             )
#         )
#     elif row_index == 2:
#         return (
#             level_3
#             if (level_1 <= level_2 and level_2 <= level_3)
#             else (
#                 level_1
#                 if (level_1 >= level_2 and level_1 >= level_3)
#                 else (
#                     level_3
#                     if (level_1 >= level_2 and level_1 < level_3)
#                     else (
#                         level_2
#                         if (level_1 <= level_2 and level_2 >= level_3)
#                         else (
#                             level_2
#                             if (level_1 <= level_2 and level_2 >= level_3)
#                             else (
#                                 level_1
#                                 if (level_1 >= level_2 and level_1 > level_3)
#                                 else level_3
#                             )
#                         )
#                     )
#                 )
#             )
#         )
